import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models import PaymentStatus
from app.schemas import WorkReport


def _money(cents: int | None) -> str:
    if cents is None:
        return "—"
    return f"${cents / 100:,.2f}"


def _dt(value) -> str:
    return value.strftime("%Y-%m-%d %H:%M") if value else "—"


def _add_heading(document: Document, text: str) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def _set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold


def _key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (key, value) in enumerate(rows):
        _set_cell_text(table.rows[i].cells[0], key, bold=True)
        _set_cell_text(table.rows[i].cells[1], value)


EVIDENCE_LABELS = {
    "route_screenshot": "Route screenshot",
    "rate_screenshot": "Rate screenshot",
    "gps_session": "GPS session",
    "completion_record": "Completion record",
    "settlement_statement": "Settlement statement",
    "other": "Other document",
}


def render_work_report_docx(report: WorkReport) -> bytes:
    document = Document()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("RINKO — INDEPENDENT WORK RECORD")
    run.bold = True
    run.font.size = Pt(15)

    ref = document.add_paragraph()
    ref.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = ref.add_run(f"Report No. {report.report_number}")
    run.font.size = Pt(10)

    document.add_paragraph()
    _key_value_table(
        document,
        [
            ("Driver:", report.driver_name),
            ("Carrier / Contractor:", report.carrier_name),
            ("Service date:", report.service_date.isoformat()),
            ("Route ID:", report.route_id or "—"),
        ],
    )

    document.add_paragraph()
    _add_heading(document, "WORK RECORD")
    wr = report.work_record
    _key_value_table(
        document,
        [
            ("Route started", _dt(wr.route_started)),
            ("Route completed", _dt(wr.route_completed)),
            ("Packages assigned", str(wr.packages_assigned)),
            ("Packages completed", str(wr.packages_completed)),
            ("Exceptions", str(wr.exceptions)),
            ("Distance", f"{wr.mileage} mi" if wr.mileage is not None else "—"),
        ],
    )

    document.add_paragraph()
    _add_heading(document, "COMPENSATION RECORD")
    comp = report.compensation
    rows = [
        ("Agreed rate", f"{_money(comp.agreed_rate_cents)}/pkg"),
        ("Expected gross", _money(comp.expected_gross_cents)),
        ("Payment due", comp.payment_due_date.isoformat() if comp.payment_due_date else "—"),
        ("Payment status", comp.payment_status.value.upper()),
    ]
    if comp.payment_status != PaymentStatus.PENDING:
        rows.append(("Payment received", _money(comp.payment_received_cents)))
        rows.append(("Received on", _dt(comp.payment_received_at)))
        if comp.difference_cents is not None and comp.difference_cents != 0:
            rows.append(("DIFFERENCE", _money(comp.difference_cents)))
    _key_value_table(document, rows)

    document.add_paragraph()
    _add_heading(document, "SUPPORTING RECORDS")
    if report.supporting_records:
        for item in report.supporting_records:
            label = EVIDENCE_LABELS.get(item.kind.value, item.kind.value)
            text = f"✓ {label}"
            if item.note:
                text += f" — {item.note}"
            document.add_paragraph(text, style="List Bullet")
    else:
        document.add_paragraph("No supporting documents attached to this session.")

    if comp.difference_cents is not None and comp.difference_cents != 0:
        document.add_paragraph()
        note = document.add_paragraph()
        note.add_run(
            f"Payment received ({_money(comp.payment_received_cents)}) differs from the expected "
            f"gross ({_money(comp.expected_gross_cents)}) by {_money(comp.difference_cents)}. "
            "This record documents the discrepancy contemporaneously; it does not by itself "
            "constitute conclusive proof in a dispute — authenticity, contract terms, and "
            "evidentiary rules still apply."
        ).italic = True

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
